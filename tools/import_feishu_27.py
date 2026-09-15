"""Import the exported 27th-cohort autumn recruitment sheet into JobPulse.

Usage:
  bundled-python tools/import_feishu_27.py <source.xlsx> [company_database.json]
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import OrderedDict
from pathlib import Path

import openpyxl


URL_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def text(value: object) -> str:
    return str(value or "").strip()


def key(value: object) -> str:
    return re.sub(r"[\s\u3000]+", "", unicodedata.normalize("NFKC", text(value)).lower())


def compact(values: list[object], limit: int = 1400) -> str:
    unique = list(OrderedDict.fromkeys(text(value) for value in values if text(value)))
    result = "；".join(unique)
    return result if len(result) <= limit else result[: limit - 1].rstrip("；") + "…"


def links(value: object, label: str) -> list[dict[str, str]]:
    # Email/application-by-email records are deliberately excluded: the public
    # library should never publish personal contact details from a source sheet.
    found: list[dict[str, str]] = []
    for _source_label, url in URL_RE.findall(text(value)):
        url = url.strip()
        if not re.match(r"https?://", url, re.I) or "@" in url:
            continue
        found.append({"text": label, "url": url})
    return found


def dedupe_links(items: list[dict[str, str]]) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    seen: set[str] = set()
    for item in items:
        url = text(item.get("url"))
        if url and url not in seen:
            seen.add(url)
            result.append({"text": text(item.get("text")) or "招聘链接", "url": url})
    return result


def make_item(name: str, records: list[tuple[object, ...]], old: dict | None) -> dict:
    industries = compact([record[2] for record in records])
    company_types = compact([record[3] for record in records])
    education = compact([record[5] for record in records])
    candidates = compact([record[6] for record in records])
    batches = compact([record[7] for record in records])
    updates = compact([record[9] for record in records])
    source_links = []
    for record in records:
        source_links.extend(links(record[13], "投递链接"))
        source_links.extend(links(record[12], "招聘公告"))

    item = dict(old or {})
    item.update(
        {
            "id": item.get("id") or "feishu27_" + hashlib.sha1(key(name).encode()).hexdigest()[:12],
            "section": industries or item.get("section") or "其他行业",
            "company": name,
            "position": compact([record[1] for record in records]),
            "location": compact([record[4] for record in records]),
            "deadline": compact([record[8] for record in records]),
            "priority": item.get("priority") or "mid",
            "priLabel": item.get("priLabel") or "⭐⭐",
            "links": dedupe_links(source_links + list(item.get("links") or [])),
            "notice": " · ".join(
                part
                for part in (
                    f"{company_types}" if company_types else "",
                    f"学历：{education}" if education else "",
                    f"招聘对象：{candidates}" if candidates else "",
                )
                if part
            ),
            "note": "；".join(
                part
                for part in (
                    "来源：飞书「27届秋招」导出",
                    f"批次：{batches}" if batches else "",
                    f"岗位更新：{updates}" if updates else "",
                )
                if part
            ),
        }
    )
    return item


def main() -> None:
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Usage: import_feishu_27.py <source.xlsx> [company_database.json]")
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2]) if len(sys.argv) == 3 else Path("company_database.json")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet.max_column < 14:
        raise ValueError("Unexpected export format: expected at least 14 columns")

    grouped: OrderedDict[str, tuple[str, list[tuple[object, ...]]]] = OrderedDict()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = text(row[0])
        if not name:
            continue
        normalized = key(name)
        if normalized not in grouped:
            grouped[normalized] = (name, [])
        grouped[normalized][1].append(row)

    database = json.loads(destination.read_text(encoding="utf-8"))
    companies = list(database.get("companies") or [])
    existing = {key(company.get("company")): company for company in companies if key(company.get("company"))}
    matched = 0
    added = 0
    updated: list[dict] = []
    for normalized, (name, records) in grouped.items():
        old = existing.get(normalized)
        if old:
            matched += 1
        else:
            added += 1
        updated.append(make_item(name, records, old))

    # Preserve the small number of unrelated library records which are absent
    # from this 27th-cohort export.
    imported_names = set(grouped)
    retained = [company for company in companies if key(company.get("company")) not in imported_names]
    database["companies"] = updated + retained
    destination.write_text(json.dumps(database, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "source_rows": sheet.max_row - 1,
        "unique_companies": len(grouped),
        "updated_existing": matched,
        "added_companies": added,
        "retained_unrelated": len(retained),
        "final_companies": len(database["companies"]),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
