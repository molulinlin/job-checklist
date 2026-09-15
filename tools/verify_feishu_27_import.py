"""Verify a 27th-cohort Feishu export was imported without field drift."""

from __future__ import annotations

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

from import_feishu_27 import compact, key, links, text


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: verify_feishu_27_import.py <source.xlsx> <company_database.json>")
    workbook = openpyxl.load_workbook(Path(sys.argv[1]), read_only=True, data_only=True)
    grouped: OrderedDict[str, tuple[str, list[tuple[object, ...]]]] = OrderedDict()
    for row in workbook.active.iter_rows(min_row=2, values_only=True):
        if text(row[0]):
            grouped.setdefault(key(row[0]), (text(row[0]), []))[1].append(row)

    database = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
    companies = database["companies"]
    by_name = {key(item["company"]): item for item in companies}
    assert len(companies) == len(grouped), (len(companies), len(grouped))
    assert len({item["id"] for item in companies}) == len(companies)
    assert all(normalized in by_name for normalized in grouped)
    assert not re.search(r"(?i)mailto:|[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}", json.dumps(companies, ensure_ascii=False))

    for normalized, (_name, records) in grouped.items():
        item = by_name[normalized]
        assert item["position"] == compact([row[1] for row in records])
        assert item["location"] == compact([row[4] for row in records])
        assert item["deadline"] == compact([row[8] for row in records])
        expected = {entry["url"] for row in records for column in (12, 13) for entry in links(row[column], "x")}
        actual = {entry["url"] for entry in item["links"]}
        assert expected <= actual, item["company"]

    print(json.dumps({"verified_companies": len(companies), "source_rows": workbook.active.max_row - 1}, ensure_ascii=False))


if __name__ == "__main__":
    main()
